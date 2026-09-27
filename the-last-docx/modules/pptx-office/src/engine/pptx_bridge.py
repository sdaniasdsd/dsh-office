"""python-pptx bridge. Input/output is JSON on stdin/stdout; diagnostics go to stderr."""
import json
import math
import os
import sys
import zipfile
from pathlib import Path


def bounded_package(path, limits):
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > limits["maxArchiveEntries"]:
            raise ValueError("PPTX archive entry count exceeds the Profile limit")
        total = 0
        for info in infos:
            if info.file_size > limits["maxEntryUncompressedBytes"]:
                raise ValueError("PPTX entry exceeds the Profile uncompressed-size limit")
            total += info.file_size
            if total > limits["maxTotalUncompressedBytes"]:
                raise ValueError("PPTX archive exceeds the Profile uncompressed-size limit")
        if "ppt/presentation.xml" not in archive.namelist():
            raise ValueError("ZIP package is not a PowerPoint presentation")
        content_types = archive.read("[Content_Types].xml")
        if b"presentationml.presentation.main+xml" not in content_types:
            raise ValueError("Only macro-free .pptx presentations are supported")
        if any(name.lower().endswith("vbaproject.bin") for name in archive.namelist()):
            raise ValueError("Macro-enabled PowerPoint files are not supported")


def shape_kind(shape):
    if getattr(shape, "has_text_frame", False):
        return "text"
    if getattr(shape, "has_table", False):
        return "table"
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "shape_type", None) is not None and str(shape.shape_type).endswith("PICTURE (13)"):
        return "picture"
    return "shape"


def extract(prs, limits):
    if len(prs.slides) > limits["maxSlides"]:
        raise ValueError("Slide count exceeds the Profile limit")
    slides = []
    shape_total = 0
    text_total = 0
    for slide_no, slide in enumerate(prs.slides, 1):
        shapes = []
        for shape in slide.shapes:
            shape_total += 1
            if shape_total > limits["maxShapes"]:
                raise ValueError("Shape count exceeds the Profile limit")
            item = {"shapeId": int(shape.shape_id), "name": str(shape.name), "kind": shape_kind(shape),
                    "bboxEmu": [int(shape.left), int(shape.top), int(shape.width), int(shape.height)]}
            if getattr(shape, "has_text_frame", False):
                paragraphs = []
                for p_no, paragraph in enumerate(shape.text_frame.paragraphs):
                    runs = []
                    for r_no, run in enumerate(paragraph.runs):
                        text = run.text or ""
                        text_total += len(text)
                        if text_total > limits["maxTextChars"]:
                            raise ValueError("PPTX text exceeds the Profile character limit")
                        font = run.font
                        size = font.size.pt if font.size is not None else None
                        runs.append({"runIndex": r_no, "text": text, "bold": font.bold, "italic": font.italic,
                                     "fontName": font.name, "fontSizePt": size})
                    paragraphs.append({"paragraphIndex": p_no, "text": paragraph.text,
                                       "alignment": str(paragraph.alignment) if paragraph.alignment is not None else None,
                                       "runs": runs})
                item["paragraphs"] = paragraphs
            elif getattr(shape, "has_table", False):
                item["rows"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
                text_total += sum(len(text) for row in item["rows"] for text in row)
                if text_total > limits["maxTextChars"]:
                    raise ValueError("PPTX text exceeds the Profile character limit")
            shapes.append(item)
        slides.append({"slideNumber": slide_no, "shapes": shapes})
    return slides


def all_text(slides):
    values = []
    for slide in slides:
        for shape in slide["shapes"]:
            for paragraph in shape.get("paragraphs", []):
                values.append(paragraph["text"])
            for row in shape.get("rows", []):
                values.extend(row)
    return "\n".join(values)


def main():
    if len(sys.argv) != 4:
        raise RuntimeError("Expected operation, input path, and output path arguments")
    operation, input_name, output_name = sys.argv[1:]
    request = json.load(sys.stdin)
    limits = request["limits"]
    bounded_package(input_name, limits)
    with open(os.devnull, "w", encoding="utf-8") as sink:
        from contextlib import redirect_stdout
        with redirect_stdout(sink):
            from pptx import Presentation
            from importlib.metadata import version
            prs = Presentation(input_name)
    if operation == "inspect":
        slides = extract(prs, limits)
        result = {"format": "pptx", "mediaType": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                  "slideCount": len(prs.slides), "shapeCount": sum(len(s["shapes"]) for s in slides),
                  "widthEmu": int(prs.slide_width), "heightEmu": int(prs.slide_height), "engine": f"python-pptx-{version('python-pptx')}"}
    elif operation == "extract":
        result = {"action": "extract", "slideCount": len(prs.slides), "slides": extract(prs, limits),
                  "engine": f"python-pptx-{version('python-pptx')}"}
    elif operation == "replaceText":
        slides = extract(prs, limits)
        changes = request.get("changes", [])
        if not changes or len(changes) > limits["maxChanges"]:
            raise ValueError("Text replacement count is empty or exceeds the Profile limit")
        seen = set()
        changed = 0
        for change in changes:
            key = (change["slideNumber"], change["shapeId"], change["paragraphIndex"], change["runIndex"])
            if key in seen:
                raise ValueError("A PPTX run may be changed only once per request")
            seen.add(key)
            slide_no, shape_id, p_no, r_no = key
            if slide_no < 1 or slide_no > len(prs.slides):
                raise ValueError("Replacement references a nonexistent slide")
            slide = prs.slides[slide_no - 1]
            shape = next((shape for shape in slide.shapes if shape.shape_id == shape_id), None)
            if shape is None or not getattr(shape, "has_text_frame", False):
                raise ValueError("Replacement target is not a text shape")
            paragraphs = shape.text_frame.paragraphs
            if p_no < 0 or p_no >= len(paragraphs) or r_no < 0 or r_no >= len(paragraphs[p_no].runs):
                raise ValueError("Replacement run pointer does not exist")
            run = paragraphs[p_no].runs[r_no]
            if run.text != change["expectedText"]:
                raise ValueError("Replacement precondition failed: the addressed run text changed")
            replacement = change["replaceWith"]
            if not isinstance(replacement, str) or "\x00" in replacement:
                raise ValueError("Replacement must be a plain text string")
            run.text = replacement
            changed += 1
        Path(output_name).parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_name)
        with redirect_stdout(sink):
            reopened = Presentation(output_name)
        new_slides = extract(reopened, limits)
        result = {"action": "replaceText", "slideCount": len(reopened.slides), "changedRuns": changed,
                  "engine": f"python-pptx-{version('python-pptx')}", "textSnapshot": all_text(new_slides)}
    elif operation == "verify":
        slides = extract(prs, limits)
        text = all_text(slides)
        checks = []
        expected_count = request.get("expectedSlideCount")
        if expected_count is not None:
            checks.append({"id": "pptx.slideCount", "status": "pass" if expected_count == len(prs.slides) else "fail",
                           "severity": "error", "message": "Reopened slide count compared with expectation."})
        for index, expected in enumerate(request.get("textIncludes", [])):
            checks.append({"id": f"pptx.text.{index}", "status": "pass" if expected in text else "fail",
                           "severity": "error", "message": "Expected text compared with extracted slide text."})
        if not checks:
            checks.append({"id": "pptx.reopened", "status": "pass", "severity": "error", "message": "PPTX reopened and structure extracted."})
        result = {"checks": checks}
    else:
        raise ValueError("Unsupported PPTX operation")
    print(json.dumps(result, ensure_ascii=False, allow_nan=False, separators=(",", ":")))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"{type(error).__name__}: {error}", file=sys.stderr)
        sys.exit(2)
