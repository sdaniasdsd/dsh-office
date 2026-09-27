# 独立核对最终 deck（不看插件的报告）：每页每个文本框的占位符类型与逐 run 字号，
# 备注是否存在，有没有漏掉没设字号的 run。
import json
import sys

from pptx import Presentation

TITLE_TYPES = {"TITLE", "CENTER_TITLE", "VERTICAL_TITLE"}
BODY_TYPES = {"BODY", "OBJECT", "SUBTITLE", "VERTICAL_BODY", "VERTICAL_OBJECT"}


def placeholder_name(shape):
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        return shape.placeholder_format.type.name          # 枚举名，例如 CENTER_TITLE
    except (AttributeError, ValueError):
        return "?"


def main():
    prs = Presentation(sys.argv[1])
    report = {"slides": len(prs.slides), "widthInches": round(prs.slide_width / 914400, 3),
              "heightInches": round(prs.slide_height / 914400, 3), "perSlide": []}
    unsized = []
    title_sizes, body_sizes = [], []
    for number, slide in enumerate(prs.slides, 1):
        shapes = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            kind = placeholder_name(shape)
            sizes, texts = [], []
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if not run.text.strip():
                        continue
                    sizes.append(None if run.font.size is None else run.font.size.pt)
                    texts.append(run.text)
                    if run.font.size is None:
                        unsized.append({"slide": number, "shapeId": int(shape.shape_id), "placeholder": kind, "text": run.text[:24]})
                    if kind in TITLE_TYPES:
                        title_sizes.append(run.font.size.pt if run.font.size else None)
                    elif kind in BODY_TYPES:
                        body_sizes.append(run.font.size.pt if run.font.size else None)
            if texts:
                shapes.append({"placeholder": kind, "sizes": sorted({s for s in sizes if s}), "text": " / ".join(texts)[:60]})
        notes = ""
        try:
            notes = slide.notes_slide.notes_text_frame.text
        except Exception:
            notes = ""
        report["perSlide"].append({"slide": number, "shapes": shapes, "notesChars": len(notes)})
    report["titleSizes"] = sorted({s for s in title_sizes if s})
    report["bodySizes"] = sorted({s for s in body_sizes if s})
    report["unsizedRuns"] = unsized
    report["totalNotesChars"] = sum(entry["notesChars"] for entry in report["perSlide"])
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
