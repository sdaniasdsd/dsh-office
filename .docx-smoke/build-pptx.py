# 用 python-pptx 建一份配套幻灯片（插件本身不含「建 deck」能力，只有 inspect / extract /
# replaceText / formatText，所以先建、再交给插件做排版层级与校验）。
#
# 关键约束（为了让插件的 formatText 能正常工作）：每一页必须用模板里的**占位符**——
# 标题占位符与正文占位符，且每个 run 都显式给出字号；否则 formatText 认不出标题/正文，
# 会直接报 “formatText found no text runs to format”。
import json
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

TITLE_COLOR = RGBColor(0x1F, 0x4E, 0x79)
BODY_COLOR = RGBColor(0x26, 0x26, 0x26)
ACCENT = RGBColor(0xC0, 0x39, 0x2B)
FONT = "Microsoft YaHei"

# 每页：标题 / 正文条目 / 备注（备注放演讲稿对应段落，讲的时候照读）
SLIDES = [
    ("嵌入式的发展路径",
     ["从一片单片机到边缘智能", "演讲人：＿＿＿＿　　日期：＿＿＿＿年＿＿月＿＿日"],
     "开场：今天我想和大家聊一个既老又新的话题——嵌入式的发展路径。说它老，是因为这门技术已经走了半个世纪；说它新，是因为最近五年，它正在被重新定义。"),
    ("今天讲三件事",
     ["嵌入式到底是什么：把计算藏进设备里",
      "它走过哪几段路：四次跃迁",
      "个人怎么走这条路：四个台阶、三条建议"],
     "我把今天的分享收在三件事上：是什么、怎么来的、个人怎么走。"),
    ("一、嵌入式是什么",
     ["不是一种芯片，也不是一种系统，而是一件事：把计算藏进设备里",
      "资源受限：内存以 KB 计、主频以 MHz 计、电量以 mAh 计",
      "行为要确定：不是“大多数时候对”，而是“每次都必须对”",
      "与物理世界绑定：传感器、执行器、时序、噪声、十年不断电",
      "判断标准：能不能在你看不见的地方长期正确运行"],
     "手机里的摄像头是嵌入式；手机上那个 App 不是。"),
    ("二、四次跃迁：总览",
     ["1980s　汇编 → C 与工具链（抽象层出现，代码第一次可读可复用）",
      "1990s–2000s　裸机 → RTOS 与 Linux（从写程序到搭系统）",
      "2010s　单机 → 联网（确定性换成分布式不确定性）",
      "2015 至今　联网 → 边缘智能（嵌入式开始承担“智能”）"],
     "回头看，这条路其实在做同一件事：在约束之下，加一层抽象。"),
    ("二·1　汇编 → C 与工具链",
     ["1980 年 Intel 8051 开启单片机时代：数机器周期、抠寄存器",
      "交叉编译器 + 链接脚本 + 调试器逐步成熟",
      "真正的成果不是语言，而是抽象层：可读、可复用、可交接"],
     "这一段的真正成果不是语言，而是抽象层：代码第一次可以被人读懂、被人复用、被人交接。"),
    ("二·2　裸机 → RTOS 与 Linux",
     ["裸机主循环 + 状态机的复杂度天花板",
      "μC/OS、2003 FreeRTOS：调度、同步、通信被标准化",
      "Linux 整包进设备：操作系统、驱动、文件系统、网络协议栈一次带齐",
      "1990 ARM 成立、2004 Cortex-M 系列：32 位单片机变便宜"],
     "工程师从“自己安排每一件事”，变成“描述任务之间的关系”。"),
    ("二·3　单机 → 联网",
     ["Wi-Fi、蓝牙低功耗、4G、窄带物联网让设备不再孤立",
      "空中升级、设备云、远程运维成为标配",
      "2005 Arduino、2012 树莓派：门槛低到学生也能做出东西",
      "代价：协议、安全、升级、故障定位全变了"],
     "联网的代价是把“单机的确定性”换成了“分布式的不确定性”，嵌入式工程师从此要背系统责任。"),
    ("二·4　联网 → 边缘智能",
     ["算力与模型下沉：模型可以跑在单片机上",
      "摄像头与麦克风在本地完成推理，采集—推理—决策闭环在现场",
      "时延、隐私、带宽三个老问题一起松动",
      "含义：嵌入式不再只是实现别人的设计，开始自己承担“智能”"],
     "这一段的含义是：嵌入式不再只是把别人的设计实现出来，它开始自己承担“智能”那部分。"),
    ("三、当下的分水岭",
     ["指令集与生态：RISC-V（2010 年起自伯克利）让“架构可以自己选”成为可讨论的题",
      "供应链与工具链自主：选型从性能优先变成性能、生态、可得性的多目标问题",
      "AI 进入开发流程本身：写功能的门槛下降，定义问题与验证可靠性的价值上升"],
     "第三个变化最容易被低估：人工智能正在进入开发流程本身。"),
    ("四、个人的路径：四个台阶",
     ["吃透一块 MCU：输入输出、中断、定时器、模数转换，加总线（I²C/SPI/串口），做出完整小项目",
      "上一门 RTOS：并发、优先级反转、栈的边界在哪里",
      "做一次真正联网的产品：能升级、能回滚、能定位",
      "选一个垂直方向扎下去：电机、电源、工业总线、汽车、端侧智能或内核"],
     "我建议按四个台阶走，不要跳。深度就是你的议价能力。"),
    ("三条建议",
     ["把“确定性地交付”练成习惯：能复现、能测量、能解释——这是嵌入式与纯软件真正的分水岭",
      "不要停在框架层：能读参考手册、看得懂时序图、会拿示波器量波形",
      "让 AI 替你干重复劳动，把省下的时间放在整机联调、边界条件与可靠性上"],
     "这些正是它暂时替不了你的地方。"),
    ("最后",
     ["这条路是“在约束之下不断加一层抽象”的历史",
      "每加一层，门槛降一级，复杂度上一级",
      "你们这一代，赶上了第四段路的开头",
      "谢谢大家！欢迎提问"],
     "我的分享就到这里，谢谢大家！欢迎提问。"),
]


def set_font(run, size, bold=False, color=BODY_COLOR):
    """同时设置拉丁字体与东亚字体：只设 latin 时，中文会回落到主题字体。"""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        for existing in rpr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1]):
            rpr.remove(existing)
        element = rpr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1], {"typeface": FONT})
        rpr.append(element)


def add_notes(slide, text):
    frame = slide.notes_slide.notes_text_frame
    frame.text = text
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(12)
            run.font.name = FONT


def main():
    target = sys.argv[1]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    cover_layout = prs.slide_layouts[0]      # Title Slide：标题 + 副标题（都是占位符）
    content_layout = prs.slide_layouts[1]    # Title and Content：标题 + 正文（都是占位符）

    for index, (title, bullets, notes) in enumerate(SLIDES):
        slide = prs.slides.add_slide(cover_layout if index == 0 else content_layout)
        title_placeholder = slide.shapes.title
        title_placeholder.left, title_placeholder.top = Inches(0.8), Inches(0.5)
        title_placeholder.width, title_placeholder.height = Inches(11.7), Inches(1.1)
        title_placeholder.text_frame.text = title
        for paragraph in title_placeholder.text_frame.paragraphs:
            for run in paragraph.runs:
                set_font(run, 36 if index == 0 else 32, bold=True, color=TITLE_COLOR if index else ACCENT)

        body = slide.placeholders[1] if index else slide.placeholders[1]
        body.left, body.top = Inches(0.9), Inches(1.9)
        body.width, body.height = Inches(11.5), Inches(4.9)
        frame = body.text_frame
        frame.word_wrap = True
        for position, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if position == 0 else frame.add_paragraph()
            paragraph.text = bullet
            paragraph.line_spacing = 1.25
            paragraph.space_after = Pt(10)
            for run in paragraph.runs:
                set_font(run, 22 if index else 24, bold=False, color=BODY_COLOR)
        add_notes(slide, notes)

    prs.save(target)
    print(json.dumps({"slides": len(prs.slides), "widthInches": round(prs.slide_width / 914400, 3),
                      "heightInches": round(prs.slide_height / 914400, 3), "target": target}, ensure_ascii=False))


if __name__ == "__main__":
    main()
