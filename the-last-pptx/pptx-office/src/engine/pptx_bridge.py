"""python-pptx bridge. Input/output is JSON on stdin/stdout; diagnostics go to stderr."""
import json
import math
import os
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path


def bounded_package(path, limits):
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > limits["maxArchiveEntries"]:
            raise ValueError("PPTX archive entry count exceeds the Profile limit")
        total = 0
        for info in infos:
            if info.file_size > limits["maxEntryUncompressedBytes"]:
                raise ValueError("PPTX entry exceeds the Profile uncompressed-size limit")
            total += info.file_size
            if total > limits["maxTotalUncompressedBytes"]:
                raise ValueError("PPTX archive exceeds the Profile uncompressed-size limit")
        if "ppt/presentation.xml" not in archive.namelist():
            raise ValueError("ZIP package is not a PowerPoint presentation")
        content_types = archive.read("[Content_Types].xml")
        if b"presentationml.presentation.main+xml" not in content_types:
            raise ValueError("Only macro-free .pptx presentations are supported")
        if any(name.lower().endswith("vbaproject.bin") for name in archive.namelist()):
            raise ValueError("Macro-enabled PowerPoint files are not supported")


def shape_kind(shape):
    if getattr(shape, "has_text_frame", False):
        return "text"
    if getattr(shape, "has_table", False):
        return "table"
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "shape_type", None) is not None and str(shape.shape_type).endswith("PICTURE (13)"):
        return "picture"
    return "shape"


def extract(prs, limits):
    if len(prs.slides) > limits["maxSlides"]:
        raise ValueError("Slide count exceeds the Profile limit")
    slides = []
    shape_total = 0
    text_total = 0
    for slide_no, slide in enumerate(prs.slides, 1):
        shapes = []
        for shape in slide.shapes:
            shape_total += 1
            if shape_total > limits["maxShapes"]:
                raise ValueError("Shape count exceeds the Profile limit")
            item = {"shapeId": int(shape.shape_id), "name": str(shape.name), "kind": shape_kind(shape),
                    "bboxEmu": [int(shape.left), int(shape.top), int(shape.width), int(shape.height)]}
            if getattr(shape, "has_text_frame", False):
                paragraphs = []
                for p_no, paragraph in enumerate(shape.text_frame.paragraphs):
                    runs = []
                    for r_no, run in enumerate(paragraph.runs):
                        text = run.text or ""
                        text_total += len(text)
                        if text_total > limits["maxTextChars"]:
                            raise ValueError("PPTX text exceeds the Profile character limit")
                        font = run.font
                        size = font.size.pt if font.size is not None else None
                        try:
                            color = font.color
                            font_color = ("#" + str(color.rgb)) if color.type is not None and color.rgb is not None else None
                        except (AttributeError, TypeError, ValueError):
                            font_color = None
                        runs.append({"runIndex": r_no, "text": text, "bold": font.bold, "italic": font.italic,
                                     "fontName": font.name, "fontSizePt": size, "fontColor": font_color})
                    paragraphs.append({"paragraphIndex": p_no, "text": paragraph.text,
                                       "alignment": str(paragraph.alignment) if paragraph.alignment is not None else None,
                                       "runs": runs})
                item["paragraphs"] = paragraphs
            elif getattr(shape, "has_table", False):
                item["rows"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
                text_total += sum(len(text) for row in item["rows"] for text in row)
                if text_total > limits["maxTextChars"]:
                    raise ValueError("PPTX text exceeds the Profile character limit")
            shapes.append(item)
        slides.append({"slideNumber": slide_no, "shapes": shapes})
    return slides


def all_text(slides):
    values = []
    for slide in slides:
        for shape in slide["shapes"]:
            for paragraph in shape.get("paragraphs", []):
                values.append(paragraph["text"])
            for row in shape.get("rows", []):
                values.extend(row)
    return "\n".join(values)


def theme_scheme_rgb(slide, scheme_name):
    master_part = slide.slide_layout.slide_master.part
    theme_rel = next((rel for rel in master_part.rels.values() if rel.reltype.endswith("/theme")), None)
    if theme_rel is None:
        return None
    root = ET.fromstring(theme_rel.target_part.blob)
    local_name = scheme_name.lower().replace("_", "")
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1] != local_name:
            continue
        color = next(iter(element), None)
        if color is None:
            return None
        value = color.attrib.get("val") or color.attrib.get("lastClr")
        if value and len(value) == 6:
            try:
                return tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))
            except ValueError:
                return None
    return None


def solid_background_rgb(slide):
    try:
        fill = slide.background.fill
        if fill.type is None or "SOLID" not in str(fill.type):
            return None
        color = fill.fore_color
        if color.type is not None and "RGB" in str(color.type):
            return tuple(color.rgb)
        if color.type is not None and "SCHEME" in str(color.type):
            theme_color = color.theme_color
            return theme_scheme_rgb(slide, getattr(theme_color, "name", str(theme_color)))
    except (AttributeError, TypeError, ValueError, KeyError):
        return None
    return None


def contrast_ratio(first, second):
    def luminance(rgb):
        channels = []
        for value in rgb:
            channel = value / 255.0
            channels.append(channel / 12.92 if channel <= 0.04045 else ((channel + 0.055) / 1.055) ** 2.4)
        return 0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2]
    high, low = sorted((luminance(first), luminance(second)), reverse=True)
    return (high + 0.05) / (low + 0.05)


def title_size_fits(shape, size_pt):
    # Conservative geometry guard: python-pptx writes direct run sizes and
    # thereby overrides PowerPoint's existing text autofit behavior. Do not
    # write a requested size when the fixed placeholder cannot contain it.
    width_pt = shape.width / 914400 * 72
    height_pt = shape.height / 914400 * 72
    text = "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs)
    explicit_lines = max(1, sum(max(1, len(part.splitlines())) for part in text.split("\v")))
    chars_per_line = max(1, int(width_pt / (size_pt * 0.55)))
    estimated_lines = max(explicit_lines, math.ceil(max(1, len(text.replace("\v", ""))) / chars_per_line))
    required_height = estimated_lines * size_pt * 1.2 + 8
    return height_pt >= required_height


def apply_indigo_paperlight_cover(prs, art_word):
    """Apply a packaged bitmap background to the cover only.

    The style registry is deliberately small and allow-listed. It keeps visual
    assets separate from the generic PPTX editing engine and avoids changing
    masters, layouts, tables, charts, or any existing text content.
    """
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.util import Inches, Pt
    asset = Path(__file__).resolve().parents[2] / "assets" / "indigo-paperlight-v1.png"
    if not asset.is_file():
        raise ValueError("Art style asset is missing from the installed module")
    if not prs.slides:
        raise ValueError("Cannot apply an art style to an empty presentation")
    slide = prs.slides[0]
    picture = slide.shapes.add_picture(str(asset), 0, 0, width=prs.slide_width, height=prs.slide_height)
    # Keep the bitmap below every existing placeholder and text shape.
    tree = slide.shapes._spTree
    tree.remove(picture._element)
    tree.insert(2, picture._element)
    watermark = slide.shapes.add_textbox(Inches(0.72), Inches(4.72), Inches(6.4), Inches(1.5))
    paragraph = watermark.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = art_word
    run.font.name = "Aptos Display"
    run.font.size = Pt(72)
    run.font.bold = True
    run.font.color.rgb = RGBColor(18, 104, 172)
    title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
    subtitle_types = {PP_PLACEHOLDER.SUBTITLE}
    styled_titles = 0
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or not getattr(shape, "is_placeholder", False):
            continue
        try:
            is_title = shape.placeholder_format.type in title_types
        except (AttributeError, ValueError):
            is_title = False
        is_subtitle = shape.placeholder_format.type in subtitle_types
        if not is_title and not is_subtitle:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if is_title:
                    run.font.name = "Aptos Display"
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
                else:
                    run.font.color.rgb = RGBColor(191, 221, 245)
        if is_title:
            styled_titles += 1
    return {"styleId": "indigo-paperlight-v1", "styledSlides": 1, "artWord": art_word,
            "styledTitlePlaceholders": styled_titles}


def apply_cool_corporate_field_cover(prs, art_word):
    """Apply an editable, restrained editorial cover composition.

    The profile follows a deliberately small design-system contract: one dark
    field, one small accent, generous light-space for the source title, and a
    single dominant native art word.  Every addition is a PowerPoint shape so
    downstream users can still edit it in PowerPoint.
    """
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
    from pptx.util import Pt
    if not prs.slides:
        raise ValueError("Cannot apply an art style to an empty presentation")
    slide = prs.slides[0]

    def field(shape, rgb, insertion_index):
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*rgb)
        shape.line.fill.background()
        tree = slide.shapes._spTree
        tree.remove(shape._element)
        tree.insert(insertion_index, shape._element)

    # The background, field and dot are kept below source content. The compact
    # dark field and small accent remain subordinate to source title content.
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    field(background, (248, 249, 250), 2)

    # Existing title/subtitle placeholders belong to the source deck.  Do not
    # rewrite their geometry: python-pptx can flatten inherited placeholder
    # bounds. Instead select decorative zones that do not overlap source text.
    def bounds(shape):
        return (shape.left / prs.slide_width, shape.top / prs.slide_height,
                shape.width / prs.slide_width, shape.height / prs.slide_height)

    def intersects(first, second):
        return not (first[0] + first[2] <= second[0] or second[0] + second[2] <= first[0]
                    or first[1] + first[3] <= second[1] or second[1] + second[3] <= first[1])

    text_bounds = [bounds(shape) for shape in slide.shapes
                   if getattr(shape, "has_text_frame", False) and shape.text.strip()]

    def free(candidates, reserved=()):
        for candidate in candidates:
            if not any(intersects(candidate, occupied) for occupied in [*text_bounds, *reserved]):
                return candidate
        raise ValueError("The cover has no collision-free zone for this design profile")

    # All coordinates are relative so the profile works on 4:3, 16:9, and
    # custom canvases. The preferred corner zone avoids common centered cover
    # placeholders; alternatives keep the profile fail-closed when it cannot.
    visual_zone = free([(0.73, 0.04, 0.23, 0.18), (0.73, 0.78, 0.23, 0.18), (0.04, 0.04, 0.23, 0.18)])
    navy_field = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(prs.slide_width * visual_zone[0]), int(prs.slide_height * visual_zone[1]), int(prs.slide_width * visual_zone[2]), int(prs.slide_height * visual_zone[3]))
    field(navy_field, (30, 58, 95), 3)
    accent_size = int(min(prs.slide_width, prs.slide_height) * 0.045)
    gold_dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(prs.slide_width * (visual_zone[0] + visual_zone[2] * 0.56)), int(prs.slide_height * (visual_zone[1] + visual_zone[3] * 0.28)), accent_size, accent_size)
    field(gold_dot, (212, 175, 55), 4)

    art_zone = free([(0.055, 0.055, 0.34, 0.12), (0.055, 0.84, 0.34, 0.10), (0.55, 0.055, 0.16, 0.12)], [visual_zone])
    art = slide.shapes.add_textbox(int(prs.slide_width * art_zone[0]), int(prs.slide_height * art_zone[1]), int(prs.slide_width * art_zone[2]), int(prs.slide_height * art_zone[3]))
    art.name = "DSH design profile art word"
    art_paragraph = art.text_frame.paragraphs[0]
    art_run = art_paragraph.add_run()
    art_run.text = art_word
    art_run.font.name = "Aptos Display"
    art_run.font.size = Pt(max(34, min(60, int(prs.slide_height / 914400 * 72 * 0.095))))
    art_run.font.bold = True
    art_run.font.color.rgb = RGBColor(30, 58, 95)

    title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
    subtitle_types = {PP_PLACEHOLDER.SUBTITLE}
    styled_titles = 0
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or not getattr(shape, "is_placeholder", False):
            continue
        try:
            placeholder_type = shape.placeholder_format.type
        except (AttributeError, ValueError):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if placeholder_type in title_types:
                    run.font.name = "Aptos Display"
                    run.font.color.rgb = RGBColor(30, 58, 95)
                    run.font.bold = True
                elif placeholder_type in subtitle_types:
                    run.font.color.rgb = RGBColor(66, 84, 102)
        if placeholder_type in title_types:
            styled_titles += 1
    return {"styleId": "cool-corporate-field-v1", "styledSlides": 1, "artWord": art_word,
            "styledTitlePlaceholders": styled_titles, "nativeObjects": 4}


def main():
    if len(sys.argv) != 4:
        raise RuntimeError("Expected operation, input path, and output path arguments")
    operation, input_name, output_name = sys.argv[1:]
    request = json.load(sys.stdin)
    limits = request["limits"]
    bounded_package(input_name, limits)
    with open(os.devnull, "w", encoding="utf-8") as sink:
        from contextlib import redirect_stdout
        with redirect_stdout(sink):
            from pptx import Presentation
            from importlib.metadata import version
            prs = Presentation(input_name)
    if operation == "inspect":
        slides = extract(prs, limits)
        result = {"format": "pptx", "mediaType": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                  "slideCount": len(prs.slides), "shapeCount": sum(len(s["shapes"]) for s in slides),
                  "widthEmu": int(prs.slide_width), "heightEmu": int(prs.slide_height), "engine": f"python-pptx-{version('python-pptx')}"}
    elif operation == "extract":
        result = {"action": "extract", "slideCount": len(prs.slides), "widthEmu": int(prs.slide_width),
                  "heightEmu": int(prs.slide_height), "slides": extract(prs, limits),
                  "engine": f"python-pptx-{version('python-pptx')}"}
    elif operation == "replaceText":
        slides = extract(prs, limits)
        changes = request.get("changes", [])
        if not changes or len(changes) > limits["maxChanges"]:
            raise ValueError("Text replacement count is empty or exceeds the Profile limit")
        seen = set()
        changed = 0
        for change in changes:
            key = (change["slideNumber"], change["shapeId"], change["paragraphIndex"], change["runIndex"])
            if key in seen:
                raise ValueError("A PPTX run may be changed only once per request")
            seen.add(key)
            slide_no, shape_id, p_no, r_no = key
            if slide_no < 1 or slide_no > len(prs.slides):
                raise ValueError("Replacement references a nonexistent slide")
            slide = prs.slides[slide_no - 1]
            shape = next((shape for shape in slide.shapes if shape.shape_id == shape_id), None)
            if shape is None or not getattr(shape, "has_text_frame", False):
                raise ValueError("Replacement target is not a text shape")
            paragraphs = shape.text_frame.paragraphs
            if p_no < 0 or p_no >= len(paragraphs) or r_no < 0 or r_no >= len(paragraphs[p_no].runs):
                raise ValueError("Replacement run pointer does not exist")
            run = paragraphs[p_no].runs[r_no]
            if run.text != change["expectedText"]:
                raise ValueError("Replacement precondition failed: the addressed run text changed")
            replacement = change["replaceWith"]
            if not isinstance(replacement, str) or "\x00" in replacement:
                raise ValueError("Replacement must be a plain text string")
            run.text = replacement
            changed += 1
        Path(output_name).parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_name)
        with redirect_stdout(sink):
            reopened = Presentation(output_name)
        new_slides = extract(reopened, limits)
        result = {"action": "replaceText", "slideCount": len(reopened.slides), "changedRuns": changed,
                  "engine": f"python-pptx-{version('python-pptx')}", "textSnapshot": all_text(new_slides)}
    elif operation == "formatText":
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import PP_PLACEHOLDER
        from pptx.util import Pt
        slides = extract(prs, limits)
        before = all_text(slides)
        changes = request.get("changes", [])
        if len(changes) != 1 or changes[0].get("scope") != "allSlides":
            raise ValueError("formatText requires one allSlides change")
        change = changes[0]
        title_size = float(change["titleFontSize"])
        body_size = float(change["bodyFontSize"])
        color_text = change["accentColor"]
        if not math.isfinite(title_size) or not math.isfinite(body_size) or not color_text.startswith("#") or len(color_text) != 7:
            raise ValueError("formatText contains invalid sizes or accentColor")
        accent = RGBColor.from_string(color_text[1:])
        formatted = 0
        formatted_slides = set()
        contrast_adjusted_slides = set()
        color_preserved_slides = set()
        expected_formats = {}
        title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
        body_types = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT}
        for slide_no, slide in enumerate(prs.slides, 1):
            background = solid_background_rgb(slide)
            title_color = None
            if background is not None:
                title_color = accent
                if contrast_ratio(tuple(accent), background) < 4.5:
                    alternatives = [(255, 255, 255), (0, 0, 0)]
                    title_color = max(alternatives, key=lambda candidate: contrast_ratio(candidate, background))
                    contrast_adjusted_slides.add(slide_no)
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False) or getattr(shape, "has_table", False):
                    continue
                placeholder_type = None
                if getattr(shape, "is_placeholder", False):
                    try:
                        placeholder_type = shape.placeholder_format.type
                    except (AttributeError, ValueError):
                        placeholder_type = None
                is_title = placeholder_type in title_types
                is_body = placeholder_type in body_types
                if not is_title and not is_body:
                    continue
                title_runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
                title_sizes = [min(title_size, run.font.size.pt) for run in title_runs if run.font.size is not None]
                fit_title_size = is_title and bool(title_runs) and len(title_sizes) == len(title_runs) and title_size_fits(shape, max(title_sizes))
                for paragraph_no, paragraph in enumerate(shape.text_frame.paragraphs):
                    for run_no, run in enumerate(paragraph.runs):
                        changed_run = False
                        expected = {}
                        if is_title:
                            if fit_title_size and run.font.size is not None:
                                wanted_size = min(title_size, run.font.size.pt)
                                run.font.size = Pt(wanted_size)
                                expected["fontSizePt"] = wanted_size
                                changed_run = True
                            if title_color is not None:
                                run.font.color.rgb = RGBColor(*title_color)
                                expected["fontColor"] = "#" + "".join(f"{value:02X}" for value in title_color)
                                changed_run = True
                            else:
                                color_preserved_slides.add(slide_no)
                        elif run.font.size is not None and run.font.size.pt > body_size:
                            # Never enlarge body copy: doing so can overrun the
                            # fixed text boxes used by dense office templates.
                            run.font.size = Pt(body_size)
                            expected["fontSizePt"] = body_size
                            changed_run = True
                        if changed_run:
                            formatted += 1
                            formatted_slides.add(slide_no)
                            expected_formats[(slide_no, int(shape.shape_id), paragraph_no, run_no)] = expected
        if formatted == 0:
            raise ValueError("formatText found no text runs to format")
        Path(output_name).parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_name)
        with redirect_stdout(sink):
            reopened = Presentation(output_name)
        new_slides = extract(reopened, limits)
        after = all_text(new_slides)
        if before != after:
            raise ValueError("formatText changed text while applying presentation formatting")
        observed = {}
        for slide in new_slides:
            for shape in slide["shapes"]:
                for paragraph in shape.get("paragraphs", []):
                    for run in paragraph.get("runs", []):
                        observed[(slide["slideNumber"], shape["shapeId"], paragraph["paragraphIndex"], run["runIndex"])] = run
        for key, expected in expected_formats.items():
            actual = observed.get(key)
            if actual is None:
                raise ValueError("formatText could not re-read a formatted run after saving")
            if "fontSizePt" in expected and (actual["fontSizePt"] is None or abs(actual["fontSizePt"] - expected["fontSizePt"]) > 0.05):
                raise ValueError("formatText font size did not survive the save/reopen check")
            if "fontColor" in expected and actual["fontColor"] != expected["fontColor"]:
                raise ValueError("formatText title color did not survive the save/reopen check")
        result = {"action": "formatText", "slideCount": len(reopened.slides), "formattedRuns": formatted,
                  "formattedSlides": len(formatted_slides), "contrastAdjustedSlides": len(contrast_adjusted_slides),
                  "titleColorPreservedSlides": len(color_preserved_slides), "verifiedFormattingRuns": len(expected_formats), "textPreserved": True,
                  "engine": f"python-pptx-{version('python-pptx')}"}
    elif operation == "applyArtStyle":
        style_id = request.get("styleId")
        styles = {
            "indigo-paperlight-v1": apply_indigo_paperlight_cover,
            "cool-corporate-field-v1": apply_cool_corporate_field_cover,
        }
        if style_id not in styles:
            raise ValueError("Unsupported art style")
        art_word = request.get("artWord", "DSH")
        if not isinstance(art_word, str) or not art_word or len(art_word) > 24:
            raise ValueError("artWord must be a short non-empty string")
        before = [paragraph["text"] for slide in extract(prs, limits) for shape in slide["shapes"]
                  for paragraph in shape.get("paragraphs", [])]
        result = styles[style_id](prs, art_word)
        Path(output_name).parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_name)
        with redirect_stdout(sink):
            reopened = Presentation(output_name)
        after = {paragraph["text"] for slide in extract(reopened, limits) for shape in slide["shapes"]
                 for paragraph in shape.get("paragraphs", [])}
        if not set(before).issubset(after):
            raise ValueError("applyArtStyle did not preserve all existing text content")
        result.update({"action": "applyArtStyle", "slideCount": len(reopened.slides), "textPreserved": True,
                       "engine": f"python-pptx-{version('python-pptx')}"})
    elif operation == "verify":
        slides = extract(prs, limits)
        text = all_text(slides)
        checks = []
        expected_count = request.get("expectedSlideCount")
        if expected_count is not None:
            checks.append({"id": "pptx.slideCount", "status": "pass" if expected_count == len(prs.slides) else "fail",
                           "severity": "error", "message": "Reopened slide count compared with expectation."})
        for index, expected in enumerate(request.get("textIncludes", [])):
            checks.append({"id": f"pptx.text.{index}", "status": "pass" if expected in text else "fail",
                           "severity": "error", "message": "Expected text compared with extracted slide text."})
        if not checks:
            checks.append({"id": "pptx.reopened", "status": "pass", "severity": "error", "message": "PPTX reopened and structure extracted."})
        result = {"checks": checks}
    else:
        raise ValueError("Unsupported PPTX operation")
    print(json.dumps(result, ensure_ascii=False, allow_nan=False, separators=(",", ":")))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"{type(error).__name__}: {error}", file=sys.stderr)
        sys.exit(2)
